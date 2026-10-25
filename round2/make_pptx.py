"""The PowerPoint file for Round Two ("A set of Power point slides of the presentation", scored on its own as "Quality
of ppt"): an editable deck made from the built HTML, so the two can never differ.

  .venv/Scripts/python.exe round2/make.py           # first: build round2/Axis.html
  .venv/Scripts/python.exe round2/make_pptx.py      # then: round2/Axis.pptx (+ PowerPoint's own render in build/pptx/)
  .venv/Scripts/python.exe round2/make_pptx.py --check round2/Axis.pptx    # print a file's metadata

How (Pete, 25 Oct: "an editable powerpoint, not photos"):
1. pptx_extract.js runs inside headless Chrome on a copy of Axis.html and reads every slide's layout: each element's
   box, fill, gradient, borders, corner radius and shadow; every text block line by line as the browser broke it,
   each run with its font, size, weight, colour and letter-spacing; pictures with their crop; each inline drawing
   (the area charts, the house, the map, icons, arrows) as standalone SVG plus a 3x PNG.
2. This script rebuilds each slide from native PowerPoint objects in the same paint order: shapes for boxes, bars,
   rules and circles; text boxes (one paragraph a line, wrap off, so the lines break where the deck's do); pictures
   for photos and portraits; drawings as SVG pictures with a PNG fallback (PowerPoint 365 can convert them to shapes).
   Speaker notes carry the spoken script (not the bracketed team notes); each slide's title is set as its name;
   sections follow the footer tracker.
3. PowerPoint itself (COM, tools/pptx_office.ps1) opens it, exports each slide as a PNG for checking, and saves a
   copy with Plus Jakarta Sans embedded, so the judges see the deck's font without installing it.
4. The saved copy's metadata is rewritten: author "Team Axis", no company, no sensitivity label, no account
   (CLAUDE.md hard rule); any e-mail address or university name left in the package stops the script.
"""
import base64
import datetime
import html as H
import io
import json
import math
import os
import re
import shutil
import subprocess
import sys
import uuid
import zipfile

from lxml import etree
from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
DECK = os.path.join(HERE, "Axis.html")
OUT = os.path.join(HERE, "Axis.pptx")                     # submitted: every slide complete, transitions only
OUT_PRESENT = os.path.join(HERE, "Axis_present.pptx")     # shared in Zoom: one click per build step
WORK = os.path.join(HERE, "build", "pptx")
CHROME = [r"C:\Program Files\Google\Chrome\Application\chrome.exe",
          r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"]
SECTIONS = [("Cover", [1]), ("The problem", [2, 3]), ("The answer", [4]), ("The five calls", [5, 6, 7, 8, 9]),
            ("Will it hold", [10]), ("What next", [11]), ("What passes on", [12]),
            ("Us", [13, 14]), ("Close", [15])]
E = 9525                                   # EMU per CSS px (a 1280 px slide is 13.333 in)
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
FONT = {400: "Plus Jakarta Sans", 500: "Plus Jakarta Sans Medium", 600: "Plus Jakarta Sans SemiBold",
        700: "Plus Jakarta Sans", 800: "Plus Jakarta Sans ExtraBold"}
plain = lambda s: re.sub(r"\s+", " ", H.unescape(re.sub(r"<[^>]+>", "", s))).strip()


# ---------------------------------------------------------------- 1. extract
def extract(deck):
    js = open(os.path.join(HERE, "pptx_extract.js"), encoding="utf-8").read()
    base = '<base href="file:///' + HERE.replace("\\", "/") + '/">'
    page = os.path.join(WORK, "_extract.html")
    with open(page, "w", encoding="utf-8") as f:
        f.write(deck.replace("<head>", "<head>" + base, 1).replace("</body>", f"<script>{js}</script></body>", 1))
    exe = next(c for c in CHROME if os.path.exists(c))
    res = subprocess.run([exe, "--headless=new", "--disable-gpu", "--window-size=1400,900", "--hide-scrollbars",
                          "--virtual-time-budget=60000", "--dump-dom", "file:///" + page.replace("\\", "/")],
                         capture_output=True, check=True)
    os.remove(page)
    dom = res.stdout.decode("utf-8")
    m = re.search(r'<pre id="pptx-data">(.*?)</pre>', dom, re.S)
    if not m:
        sys.exit("extraction failed: no data in the page")
    data = json.loads(H.unescape(m.group(1)))
    if "error" in data:
        sys.exit("extraction failed: " + data["error"])
    return data


def slide_info(deck):
    """(title, spoken script) for each slide, from the built HTML."""
    out = []
    for body in re.findall(r'<section class="slide[^"]*"[^>]*>(.*?)</section>', deck, re.S):
        t = re.search(r"<h2>(.*?)</h2>", body, re.S) or re.search(r"<h1>(.*?)</h1>", body, re.S)
        title = plain(t.group(1).replace("<br>", " ")) if t else ""
        n = re.search(r'<aside class="notes">(.*?)</aside>', body, re.S)
        spoken = plain(n.group(1)).split("[Notes:")[0].strip() if n else ""
        parts = re.split(r"\b(PETE|FAHTAI|WIN|LOOKBUA|ALL):\s*", spoken)
        lines = [f"{who.title()}: {txt.strip()}" for who, txt in zip(parts[1::2], parts[2::2]) if txt.strip()]
        out.append((title, "\n".join(lines) or spoken))
    return out


# ---------------------------------------------------------------- 2. build: helpers
def emu(v):
    return Emu(int(round(v * E)))


def a(tag, **attrs):
    el = etree.Element(qn("a:" + tag))
    for k, v in attrs.items():
        el.set(k, str(v))
    return el


def clr(c, op=1.0):
    """an <a:srgbClr> from [r, g, b, alpha] and an extra opacity"""
    el = a("srgbClr", val="%02X%02X%02X" % tuple(int(round(x)) for x in c[:3]))
    al = (c[3] if len(c) > 3 else 1) * op
    if al < 0.999:
        el.append(a("alpha", val=int(round(al * 100000))))
    return el


def split_top(s):
    out, depth, cur = [], 0, ""
    for ch in s:
        if ch == "(":
            depth += 1
        elif ch == ")":
            depth -= 1
        if ch == "," and depth == 0:
            out.append(cur.strip())
            cur = ""
        else:
            cur += ch
    if cur.strip():
        out.append(cur.strip())
    return out


def parse_rgba(s):
    m = re.search(r"rgba?\(([^)]+)\)", s)
    if not m:
        return None
    p = [float(x) for x in re.split(r"[\s,/]+", m.group(1).strip()) if x]
    return p[:3] + [p[3] if len(p) > 3 else 1]


def grad_fill(css, w, h, op):
    """a CSS linear/radial gradient (as computed) -> <a:gradFill>"""
    layer = split_top(css)[0]
    m = re.match(r"(?:repeating-)?(linear|radial)-gradient\((.*)\)$", layer.strip(), re.S)
    if not m:
        return None
    kind, args = m.group(1), split_top(m.group(2))
    ang = 180.0
    if args and not args[0].startswith("rgb"):
        head = args.pop(0)
        dm = re.match(r"(-?[\d.]+)deg", head)
        if dm:
            ang = float(dm.group(1))
        elif head.startswith("to "):
            ang = {"to top": 0, "to right": 90, "to bottom": 180, "to left": 270}.get(head, 180)
    length = abs(w * math.sin(math.radians(ang))) + abs(h * math.cos(math.radians(ang))) or 1
    stops = []
    for i, s in enumerate(args):
        c = parse_rgba(s)
        if not c:
            continue
        rest = s[s.index(")") + 1:].strip()
        if rest.endswith("%"):
            pos = float(rest[:-1]) / 100
        elif rest.endswith("px"):
            pos = float(rest[:-2]) / length
        else:
            pos = None
        stops.append([pos, c])
    if len(stops) < 2:
        return None
    for i, s in enumerate(stops):
        if s[0] is None:
            s[0] = i / (len(stops) - 1)
    g = a("gradFill", rotWithShape="0")                    # CSS gradients follow the page, not a rotated shape
    gl = a("gsLst")
    for pos, c in stops:
        gs = a("gs", pos=int(round(max(0, min(1, pos)) * 100000)))
        gs.append(clr(c, op))
        gl.append(gs)
    g.append(gl)
    if kind == "linear":
        g.append(a("lin", ang=int(round(((ang - 90) % 360) * 60000)), scaled="0"))
    else:
        p = a("path", path="circle")
        p.append(a("fillToRect", l="50000", t="50000", r="50000", b="50000"))
        g.append(p)
    return g


def parse_shadow(css):
    for s in split_top(css):
        if "inset" in s:
            continue
        c = parse_rgba(s) or [0, 0, 0, 0.3]
        nums = [float(x) for x in re.findall(r"(-?[\d.]+)px", re.sub(r"rgba?\([^)]*\)", "", s))]
        nums += [0] * (4 - len(nums))
        return {"c": c, "dx": nums[0], "dy": nums[1], "blur": nums[2], "spread": nums[3]}
    return None


def shadow_el(sh, op):
    eff = a("effectLst")
    dist = math.hypot(sh["dx"], sh["dy"])
    o = a("outerShdw", blurRad=int(sh["blur"] * E), dist=int(dist * E),
          dir=int((math.degrees(math.atan2(sh["dy"], sh["dx"])) % 360) * 60000), algn="ctr", rotWithShape="0")
    o.append(clr(sh["c"], op))
    eff.append(o)
    return eff


def strip_style(shape):
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def set_paint(shape, fill=None, line=None, effect=None):
    """write fill, outline and effects into a shape's spPr, in schema order, replacing what python-pptx wrote"""
    sp = shape._element.spPr
    for tag in ("a:noFill", "a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill", "a:grpFill", "a:ln", "a:effectLst"):
        for el in sp.findall(qn(tag)):
            sp.remove(el)
    geom = sp.find(qn("a:prstGeom"))
    if geom is None:
        geom = sp.find(qn("a:custGeom"))
    at = list(sp).index(geom) + 1 if geom is not None else len(sp)
    seq = [fill if fill is not None else a("noFill")]
    if line is not None:
        seq.append(line)
    else:
        ln = a("ln")
        ln.append(a("noFill"))
        seq.append(ln)
    if effect is not None:
        seq.append(effect)
    for i, el in enumerate(seq):
        sp.insert(at + i, el)


def solid(c, op=1.0):
    f = a("solidFill")
    f.append(clr(c, op))
    return f


def outline(w, c, op, dash=None):
    ln = a("ln", w=int(round(w * E)))
    ln.append(solid(c, op))
    if dash:
        ln.append(a("prstDash", val=dash))
    return ln


def geometry(rad, w, h):
    """(MSO_SHAPE, rotation, adj) for CSS corner radii [tl, tr, br, bl]"""
    m = min(w, h)
    if m <= 0 or max(rad) < 0.5:
        return MSO_SHAPE.RECTANGLE, 0, None
    tl, tr, br, bl = [min(r, m / 2) for r in rad]
    if min(tl, tr, br, bl) >= m / 2 - 0.6:
        return (MSO_SHAPE.OVAL, 0, None) if abs(w - h) < 2 else (MSO_SHAPE.ROUNDED_RECTANGLE, 0, 0.5)
    if abs(tl - tr) < 0.6 and abs(tl - br) < 0.6 and abs(tl - bl) < 0.6:
        return MSO_SHAPE.ROUNDED_RECTANGLE, 0, tl / m
    if tl > 0.5 and abs(tl - bl) < 0.6 and tr < 0.5 and br < 0.5:
        return MSO_SHAPE.ROUND_2_SAME_RECTANGLE, 270, tl / m
    if tr > 0.5 and abs(tr - br) < 0.6 and tl < 0.5 and bl < 0.5:
        return MSO_SHAPE.ROUND_2_SAME_RECTANGLE, 90, tr / m
    if tl > 0.5 and abs(tl - tr) < 0.6 and bl < 0.5 and br < 0.5:
        return MSO_SHAPE.ROUND_2_SAME_RECTANGLE, 0, tl / m
    if bl > 0.5 and abs(bl - br) < 0.6 and tl < 0.5 and tr < 0.5:
        return MSO_SHAPE.ROUND_2_SAME_RECTANGLE, 180, bl / m
    return MSO_SHAPE.ROUNDED_RECTANGLE, 0, sum(rad) / 4 / m


def add_geom(slide, kind, rot, adj, x, y, w, h):
    if rot in (90, 270):
        cx, cy = x + w / 2, y + h / 2
        x, y, w, h = cx - h / 2, cy - w / 2, h, w
    shp = slide.shapes.add_shape(kind, emu(x), emu(y), emu(max(w, 0.5)), emu(max(h, 0.5)))
    strip_style(shp)
    if rot:
        shp.rotation = rot
    if adj is not None and len(shp.adjustments):
        shp.adjustments[0] = max(0.0, min(0.5, adj))
    return shp


# ---------------------------------------------------------------- 2. build: items
def add_box(slide, it):
    x, y, w, h, op = it["x"], it["y"], it["w"], it["h"], it["op"]
    if w < 0.3 or h < 0.3:
        return
    if it.get("tri"):                                       # a CSS border triangle -> a real triangle, pointing away
        rot = {"b": 0, "l": 90, "t": 180, "r": 270}[it["tri"]["side"]]
        shp = add_geom(slide, MSO_SHAPE.ISOSCELES_TRIANGLE, rot, None, x, y, w, h)
        set_paint(shp, fill=solid(it["tri"]["c"], op))
        shp.name = "Arrowhead"
        return
    kind, rot, adj = geometry(it.get("rad", [0, 0, 0, 0]), w, h)
    border = it.get("border") or {}
    same = len(border) == 4 and len({(b["w"], tuple(b["c"]), b["st"]) for b in border.values()}) == 1
    sh = parse_shadow(it["shadow"]) if it.get("shadow") else None
    if sh and sh["blur"] == 0 and sh["dx"] == 0 and sh["dy"] == 0 and sh["spread"] > 0:
        s = sh["spread"]                                    # a spread-only shadow is a ring: draw it behind
        ring = add_geom(slide, kind, rot, adj if adj is None else adj, x - s, y - s, w + 2 * s, h + 2 * s)
        set_paint(ring, fill=solid(sh["c"], op))
        ring.name = "Ring"
        sh = None
    fill = None
    if it.get("grad"):
        fill = grad_fill(it["grad"], w, h, op)
    if fill is None and it.get("fill"):
        fill = solid(it["fill"], op)
    line = None
    if same:
        b = next(iter(border.values()))
        dash = {"dashed": "dash", "dotted": "sysDot"}.get(b["st"])
        line = outline(b["w"], b["c"], op, dash)
        x, y, w, h = x + b["w"] / 2, y + b["w"] / 2, w - b["w"], h - b["w"]
    if fill is not None or line is not None or sh:
        shp = add_geom(slide, kind, rot, adj, x, y, w, h)
        set_paint(shp, fill=fill, line=line, effect=shadow_el(sh, op) if sh else None)
        shp.name = it.get("tag", "Box")
    if border and not same:
        X, Y, W, Hh = it["x"], it["y"], it["w"], it["h"]
        for side, b in border.items():
            bw = b["w"]
            if b["st"] in ("dashed", "dotted"):
                dash = "dash" if b["st"] == "dashed" else "sysDot"
                p = {"t": (X, Y + bw / 2, X + W, Y + bw / 2), "b": (X, Y + Hh - bw / 2, X + W, Y + Hh - bw / 2),
                     "l": (X + bw / 2, Y, X + bw / 2, Y + Hh), "r": (X + W - bw / 2, Y, X + W - bw / 2, Y + Hh)}[side]
                ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, emu(p[0]), emu(p[1]), emu(p[2]), emu(p[3]))
                ln.line.width = emu(bw)
                ln.line.color.rgb = __import__("pptx").dml.color.RGBColor(*[int(v) for v in b["c"][:3]])
                ln.line.dash_style = {"dash": 4, "sysDot": 3}[dash]
                ln.name = "Rule"
            else:
                r = {"t": (X, Y, W, bw), "b": (X, Y + Hh - bw, W, bw), "l": (X, Y, bw, Hh), "r": (X + W - bw, Y, bw, Hh)}[side]
                shp = add_geom(slide, MSO_SHAPE.RECTANGLE, 0, None, *r)
                set_paint(shp, fill=solid(b["c"], op))
                shp.name = "Rule"


def add_text(slide, it, cal):
    lines = it["lines"]
    align = it["align"]
    lefts = [L["l"] for L in lines]
    rights = [L["r"] for L in lines]
    top = lines[0]["mid"] - lines[0]["lh"] / 2
    bottom = lines[-1]["mid"] + lines[-1]["lh"] / 2
    if align in ("center", "-webkit-center"):
        cx = (min(lefts) + max(rights)) / 2
        w = (max(rights) - min(lefts)) * 1.06 + 12
        x, al = cx - w / 2, PP_ALIGN.CENTER
    elif align in ("right", "end", "-webkit-right"):
        w = (max(rights) - min(lefts)) * 1.06 + 12
        x, al = max(rights) - w, PP_ALIGN.RIGHT
    else:
        x, w, al = min(lefts), (max(rights) - min(lefts)) * 1.06 + 12, PP_ALIGN.LEFT
    tb = slide.shapes.add_textbox(emu(x), emu(top), emu(w), emu(max(bottom - top, 4)))
    tb.name = "Text"
    tf = tb.text_frame
    tf.word_wrap = False
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    prev = None
    for i, L in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = al
        p.line_spacing = Pt(L["lh"] * 0.75)
        gap = 0 if prev is None else (L["mid"] - prev["mid"]) - (prev["lh"] + L["lh"]) / 2
        p.space_before = Pt(max(0, gap) * 0.75)
        p.space_after = Pt(0)
        runs = []
        last = None
        for r in L["runs"]:                                # merge neighbours with the same look
            r = dict(r)
            if last is not None:
                # a margin or an inline picture (a swatch) between two runs: spaces of the same width (the font's
                # space is 0.17 em), so the words after it stay where the deck has them
                gap = r["l"] - last["r"]
                if gap > 0.15 * r["fs"] and not last["s"].endswith(" ") and not r["s"].startswith(" "):
                    r["s"] = " " * max(1, round(gap / (0.17 * r["fs"]))) + r["s"]
                elif gap > 0.5 * r["fs"]:
                    r["s"] = " " * max(1, round((gap - 0.17 * r["fs"]) / (0.17 * r["fs"]))) + r["s"]
            last = r
            key = (r["fs"], r["fw"], r["it"], tuple(r["col"] or []), r["ls"], round(r["op"], 3), r.get("u"))
            if runs and runs[-1][0] == key:
                runs[-1][1]["s"] += r["s"]
            else:
                runs.append([key, r])
        for _, r in runs:
            run = p.add_run()
            run.text = r["s"]
            f = run.font
            fw = int(r["fw"])
            fw = 400 if fw < 450 else 500 if fw < 550 else 600 if fw < 650 else 700 if fw < 750 else 800
            f.name = FONT[fw]
            f.size = Pt(r["fs"] * 0.75)
            f.bold = fw == 700
            f.italic = bool(r["it"])
            if r.get("u"):
                f.underline = True
            rPr = run._r.get_or_add_rPr()
            rPr.set("kern", "100")
            if r["ls"]:
                rPr.set("spc", str(int(round(r["ls"] * 75))))
            if r["col"]:
                for old in rPr.findall(qn("a:solidFill")):
                    rPr.remove(old)
                rPr.insert(0, solid(r["col"], r["op"]))
            for tag in ("a:latin", "a:ea", "a:cs"):
                el = rPr.find(qn(tag))
                if el is not None:
                    el.set("typeface", FONT[fw])
        prev = L
    # PowerPoint sets a line's text lower in an exact-spaced line than the browser does: nudge the box up by a
    # share of the first line's largest font size (measured against PowerPoint's own render, see --calibrate)
    big = max(r["fs"] for r in lines[0]["runs"])
    tb.top = emu(top - cal * big)


def add_picture(slide, prs, it, assets, svg_parts):
    ast = assets[it["a"]]
    x, y, w, h = it["x"], it["y"], it["w"], it["h"]
    if w < 0.5 or h < 0.5:
        return
    if ast["kind"] in ("svg", "svgimg"):
        if not ast.get("png"):
            return
        blob = base64.b64decode(ast["png"].split(",", 1)[1])
    else:
        blob = base64.b64decode(ast["src"].split(",", 1)[1])
    crop = [0.0, 0.0, 0.0, 0.0]                             # l, t, r, b
    if it.get("k") == "img" and it.get("fit") == "cover" and ast.get("nw"):
        nw, nh = ast["nw"], ast["nh"]
        s = max(w / nw, h / nh)
        px, py = [float(v[:-1]) / 100 if v.endswith("%") else 0.5 for v in (it.get("pos") or "50% 50%").split()[:2]]
        ox, oy = (w - nw * s) * px, (h - nh * s) * py
        l, t = -ox / s, -oy / s
        crop = [l / nw, t / nh, 1 - (l + w / s) / nw, 1 - (t + h / s) / nh]
    clip = it.get("clip")
    round_ = bool(it.get("round"))
    if clip:                                                # crop to the parent that clips it
        cx0, cy0 = max(x, clip["x"]), max(y, clip["y"])
        cx1, cy1 = min(x + w, clip["x"] + clip["w"]), min(y + h, clip["y"] + clip["h"])
        if cx1 > cx0 and cy1 > cy0 and (cx0 > x + 0.5 or cy0 > y + 0.5 or cx1 < x + w - 0.5 or cy1 < y + h - 0.5):
            vw, vh = 1 - crop[0] - crop[2], 1 - crop[1] - crop[3]
            crop = [crop[0] + (cx0 - x) / w * vw, crop[1] + (cy0 - y) / h * vh,
                    crop[2] + (x + w - cx1) / w * vw, crop[3] + (y + h - cy1) / h * vh]
            x, y, w, h = cx0, cy0, cx1 - cx0, cy1 - cy0
        round_ = round_ or clip["round"]
    pic = slide.shapes.add_picture(io.BytesIO(blob), emu(x), emu(y), emu(w), emu(h))
    pic.name = "Drawing" if ast["kind"] in ("svg", "svgimg") else "Picture"
    if any(abs(c) > 1e-4 for c in crop):
        pic.crop_left, pic.crop_top, pic.crop_right, pic.crop_bottom = crop
    if round_:
        pic._element.spPr.find(qn("a:prstGeom")).set("prst", "ellipse")
    if it.get("border") and len(it["border"]) == 4:
        b = next(iter(it["border"].values()))
        pic.line.width = emu(b["w"])
        pic.line.color.rgb = __import__("pptx").dml.color.RGBColor(*[int(v) for v in b["c"][:3]])
    if it.get("shadow"):
        sh = parse_shadow(it["shadow"])
        if sh and sh["blur"] > 0:
            sp = pic._element.spPr
            sp.append(shadow_el(sh, it["op"]))
    if it["op"] < 0.98:
        blip = pic._element.blipFill.find(qn("a:blip"))
        blip.append(a("alphaModFix", amt=int(it["op"] * 100000)))
    if it.get("grey"):
        pic._element.blipFill.find(qn("a:blip")).append(a("grayscl"))
    if ast["kind"] in ("svg", "svgimg") and ast.get("svg"):    # the vector original, PowerPoint 365 shows it
        key = it["a"]
        if key not in svg_parts:
            n = len(svg_parts) + 1
            svg_parts[key] = Part(PackURI(f"/ppt/media/drawing{n}.svg"), "image/svg+xml", prs.part.package,
                                  ast["svg"].encode("utf-8"))
        rid = pic.part.relate_to(svg_parts[key], RT.IMAGE)
        blip = pic._element.blipFill.find(qn("a:blip"))
        ext_lst = blip.find(qn("a:extLst"))
        if ext_lst is None:
            ext_lst = etree.SubElement(blip, qn("a:extLst"))
        ext = etree.SubElement(ext_lst, qn("a:ext"))
        ext.set("uri", "{96DAC541-7B7A-43D3-8B79-37D633B846F1}")
        sb = etree.SubElement(ext, "{http://schemas.microsoft.com/office/drawing/2016/SVG/main}svgBlip",
                              nsmap={"asvg": "http://schemas.microsoft.com/office/drawing/2016/SVG/main"})
        sb.set(qn("r:embed"), rid)


def add_sections(prs):
    ids = [s.slide_id for s in prs.slides]
    P = "http://schemas.openxmlformats.org/presentationml/2006/main"
    P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
    secs = "".join(
        f'<p14:section name="{H.escape(name)}" id="{{{str(uuid.uuid4()).upper()}}}"><p14:sldIdLst>'
        + "".join(f'<p14:sldId id="{ids[i - 1]}"/>' for i in nums) + "</p14:sldIdLst></p14:section>"
        for name, nums in SECTIONS)
    ext = etree.fromstring(f'<p:ext xmlns:p="{P}" uri="{{521415D9-36F7-43E2-AB2F-B90AF26B5E84}}">'
                           f'<p14:sectionLst xmlns:p14="{P14}">{secs}</p14:sectionLst></p:ext>')
    root = prs.part._element
    lst = root.find(f"{{{P}}}extLst")
    if lst is None:
        lst = etree.SubElement(root, f"{{{P}}}extLst")
    lst.append(ext)


# ---------------------------------------------------------------- 2b. motion: transitions, builds
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"


def grey_rgba(c):
    g = 0.299 * c[0] + 0.587 * c[1] + 0.114 * c[2]
    return [g, g, g] + list(c[3:4] or [1])


def greyed(it):
    """the deck's 'still to come' look (filter: grayscale(1); opacity: .3), for slides 2, 4 and 14"""
    import copy
    it = copy.deepcopy(it)
    it["op"] = it.get("op", 1) * 0.3
    if it.get("fill"):
        it["fill"] = grey_rgba(it["fill"])
    if it.get("tri"):
        it["tri"]["c"] = grey_rgba(it["tri"]["c"])
    for b in (it.get("border") or {}).values():
        b["c"] = grey_rgba(b["c"])
    sub = lambda s: re.sub(r"rgba?\(([^)]+)\)", lambda m: "rgba(%s)" % ", ".join(
        [str(round(grey_rgba(parse_rgba(m.group(0)))[0]))] * 3 + [str(parse_rgba(m.group(0))[3])]), s)
    for key in ("grad", "shadow"):
        if it.get(key):
            it[key] = sub(it[key])
    for L in it.get("lines", []):
        for r in L["runs"]:
            r["op"] = r.get("op", 1) * 0.3
            if r.get("col"):
                r["col"] = grey_rgba(r["col"])
    if it["k"] in ("img", "svg"):
        it["grey"] = True
    return it


def transition_el(kind):
    """Morph into the slides that keep the same frame and house (5 to 9); no transition elsewhere (Pete, 25 Oct:
    no fade between slides). Returns None for no transition."""
    if kind != "morph":
        return None
    if kind == "morph":
        xml = ('<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
               '<mc:Choice xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main" Requires="p159">'
               f'<p:transition xmlns:p="{P_NS}" xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" '
               'spd="slow" p14:dur="900"><p159:morph option="byObject"/></p:transition></mc:Choice>'
               f'<mc:Fallback><p:transition xmlns:p="{P_NS}" spd="fast"><p:fade/></p:transition></mc:Fallback>'
               '</mc:AlternateContent>')
    else:
        xml = f'<p:transition xmlns:p="{P_NS}" spd="fast"><p:fade/></p:transition>'
    return etree.fromstring(xml)


def timing_el(clicks, bld):
    """the slide's main sequence: one click per build step; each shape fades in (300 ms), all of a step together;
    'exits' disappear in the same click (the house swap before the Morph). bld: spids of shapes with a text body."""
    n = [2]

    def nid():
        n[0] += 1
        return n[0]

    def effect(spid, first, exit_=False):
        a_id, s_id = nid(), nid()
        if exit_:                                       # after the entrances' 300 ms fade, so nothing blinks
            return (f'<p:par><p:cTn id="{a_id}" presetID="1" presetClass="exit" presetSubtype="0" fill="hold" grpId="1" '
                    f'nodeType="{"clickEffect" if first else "withEffect"}"><p:stCondLst><p:cond delay="320"/></p:stCondLst>'
                    f'<p:childTnLst><p:set><p:cBhvr><p:cTn id="{s_id}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/>'
                    f'</p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>'
                    'style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="hidden"/></p:to></p:set>'
                    '</p:childTnLst></p:cTn></p:par>')
        f_id = nid()
        return (f'<p:par><p:cTn id="{a_id}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" '
                f'nodeType="{"clickEffect" if first else "withEffect"}"><p:stCondLst><p:cond delay="0"/></p:stCondLst>'
                f'<p:childTnLst><p:set><p:cBhvr><p:cTn id="{s_id}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/>'
                f'</p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl><p:attrNameLst><p:attrName>'
                'style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
                f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{f_id}" dur="300"/><p:tgtEl>'
                f'<p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:animEffect></p:childTnLst></p:cTn></p:par>')

    groups = ""
    for ins, outs in clicks:
        g1, g2 = nid(), nid()
        eff = "".join(effect(sp, i == 0) for i, sp in enumerate(ins))
        eff += "".join(effect(sp, not ins and i == 0, exit_=True) for i, sp in enumerate(outs))
        groups += (f'<p:par><p:cTn id="{g1}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst>'
                   f'<p:childTnLst><p:par><p:cTn id="{g2}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst>'
                   f'<p:childTnLst>{eff}</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>')
    blds = "".join(f'<p:bldP spid="{sp}" grpId="{g}" animBg="1"/>' for sp, g in bld)
    xml = (f'<p:timing xmlns:p="{P_NS}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">'
           '<p:childTnLst><p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq">'
           f'<p:childTnLst>{groups}</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl>'
           '<p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl>'
           '<p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq></p:childTnLst></p:cTn></p:par></p:tnLst>'
           + (f'<p:bldLst>{blds}</p:bldLst>' if blds else "") + '</p:timing>')
    return etree.fromstring(xml)


def first_text(items):
    def line(L):
        out, last = "", None
        for r in L["runs"]:
            if last is not None and r["l"] - last["r"] > 0.15 * r["fs"] and not out.endswith(" "):
                out += " "
            out += r["s"]
            last = r
        return out.strip()
    for it in items:
        if it["k"] == "text":
            t = " ".join(line(L) for L in it["lines"]).strip()
            if t:
                return t if len(t) <= 60 else t[:57] + "…"
    return "(a drawing)"


def build(data, info, path, cal, present=False):
    """present=False: the submitted file, every slide complete. present=True: the copy shared in Zoom, each build
    step one click (the HTML deck's data-b order, which the script follows)."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(1280 * E), Emu(720 * E)
    blank = prs.slide_layouts[6]
    svg_parts = {}
    counts = []
    has_house = [any(it.get("role") == "house" for it in sl["items"]) for sl in data["slides"]]
    for i, (sl, (title, spoken)) in enumerate(zip(data["slides"], info), 1):
        s = prs.slides.add_slide(blank)
        items = sl["items"]
        if present:
            items = [it for it in items if not it.get("whole") or it.get("role")]
        else:
            items = [it for it in items if not it.get("layer")]
        # paint order: the deck's stacking (z-index), stable, the slide's own background first
        items = sorted(items, key=lambda it: -10 ** 9 if it.get("tag") == "section" else it.get("z", 0))
        last = max([it.get("b", 1) for it in items] + [1])
        steps, swap_in, swap_out = {}, [], []
        for it in items:
            k = it["k"]
            if it.get("tag") == "section" and (it.get("grad") or it.get("fill")):
                # the slide's own colour (slide 4's wine stage) as the slide background, not a shape: Morph
                # cross-fades backgrounds, but would fly a full-slide shape into some shape on the next slide
                fill = grad_fill(it["grad"], it["w"], it["h"], 1) if it.get("grad") else None
                if fill is None:
                    fill = solid(it["fill"])
                bg = etree.Element(qn("p:bg"))
                bgpr = etree.SubElement(bg, qn("p:bgPr"))
                bgpr.append(fill)
                bgpr.append(a("effectLst"))
                s._element.cSld.insert(0, bg)
                continue
            draw = add_box if k == "box" else (lambda sl_, it_: add_text(sl_, it_, cal)) if k == "text" else \
                (lambda sl_, it_: add_picture(sl_, prs, it_, data["assets"], svg_parts))
            step = it.get("b", 1) if present else 1
            if present and it.get("whole"):                 # the whole house, swapped in on the last click for Morph
                n0 = len(s.shapes)
                draw(s, it)
                swap_in += [sh.shape_id for sh in list(s.shapes)[n0:]]
                continue
            if step > 1 and not sl["reveal"]:              # still to come: grey until spoken
                n0 = len(s.shapes)
                draw(s, greyed(it))
                if it.get("role") and it.get("layer"):
                    swap_out += [sh.shape_id for sh in list(s.shapes)[n0:]]
            n0 = len(s.shapes)
            draw(s, it)
            new = list(s.shapes)[n0:]
            if step > 1:
                steps.setdefault(step, []).extend(sh.shape_id for sh in new)
            if it.get("role") and it.get("layer"):
                swap_out += [sh.shape_id for sh in new]
            elif it.get("role") == "house":                  # one name on every slide, so Morph carries the house
                for sh in new:
                    sh.name = "!!house"
        for sh in s.shapes:
            if sh.shape_id in swap_in:
                sh.name = "!!house"
        # names for Morph: identical things on two slides (the footer's tracker, the page frame) share a name and stay
        # put; everything else gets a name of its own, so Morph fades it instead of flying unrelated text across
        import hashlib
        seen = {}
        for sh in s.shapes:
            if sh.name == "!!house":
                continue
            if sh.has_text_frame and sh.text_frame.text.strip():
                sig = "T " + re.sub(r"\s+", " ", sh.text_frame.text.strip())[:40]
            else:
                sig = sh.name + " " + hashlib.md5(etree.tostring(sh._element.spPr if hasattr(sh._element, "spPr")
                                                                 else sh._element)).hexdigest()[:8]
            seen[sig] = seen.get(sig, 0) + 1
            # "!!" tells Morph to pair a shape only with the same name on the next slide; without it, Morph also
            # pairs unrelated shapes of one kind (slide 4's result box flew into slide 5's rules, 25 Oct)
            sh.name = "!!" + (sig if seen[sig] == 1 else f"{sig} #{seen[sig]}")
        # transition into this slide
        kind = "morph" if i > 1 and has_house[i - 1] and has_house[i - 2] else "fade"
        root = s._element
        tr = transition_el(kind)
        if tr is not None:
            root.insert(list(root).index(root.find(qn("p:clrMapOvr"))) + 1, tr)
        if present and (steps or swap_in):
            clicks = [(steps[k], []) for k in sorted(steps)]
            if swap_in:
                if clicks:
                    clicks[-1] = (clicks[-1][0] + swap_in, swap_out)
                else:
                    clicks = [(swap_in, swap_out)]
            sp_ids = {int(sp.get("id")) for sp in root.iter(qn("p:cNvPr")) if sp.getparent().getparent().tag == qn("p:sp")}
            bld = [(sp, 0) for sp in sorted({x for c in clicks for x in c[0]} & sp_ids)]
            bld += [(sp, 1) for sp in sorted({x for c in clicks for x in c[1]} & sp_ids)]
            root.append(timing_el(clicks, bld))
            cues = [f"Click {n}: {first_text([it for it in items if it.get('b', 1) == k and not it.get('whole')])}"
                    for n, k in enumerate(sorted(steps), 1)]
            spoken = spoken + "\n\n[Clicks on this slide]\n" + "\n".join(cues)
        s.notes_slide.notes_text_frame.text = spoken
        s._element.cSld.set("name", title[:250])
        counts.append(len(s.shapes))
    add_sections(prs)
    # embed whole fonts, not the characters used: a teammate without Plus Jakarta Sans can still type new words
    prs.part._element.attrib.pop("saveSubsetFonts", None)
    cp = prs.core_properties
    now = datetime.datetime.now(datetime.timezone.utc).replace(tzinfo=None, microsecond=0)
    cp.author = cp.last_modified_by = "Team Axis"
    cp.title, cp.subject = "Built to Last", "UFP Award 2026, Round Two: a financial proposal for the Wong family"
    cp.keywords = cp.comments = cp.category = ""
    cp.revision, cp.created, cp.modified = 1, now, now
    prs.save(path)
    return counts


# ---------------------------------------------------------------- 3. PowerPoint: render, embed the font
def office(src, out, pngdir):
    ps = os.path.join(ROOT, "tools", "pptx_office.ps1")
    args = ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", ps, "-In", src]
    if pngdir:
        args += ["-PngDir", pngdir]
    if out:
        args += ["-Out", out]
    r = subprocess.run(args, capture_output=True, text=True)
    if r.returncode:
        sys.exit("PowerPoint step failed:\n" + r.stdout + r.stderr)
    return r.stdout.strip()


# ---------------------------------------------------------------- 4. metadata
DROP = re.compile(r"^(docMetadata/|customXml/|docProps/custom\.xml|ppt/authors\.xml|ppt/commentAuthors\.xml)")
BAD = re.compile(r"[\w.+-]+@[\w-]+\.[a-z.]{2,}|\\Users\\|/Users/|peeradon|zushi|\b(hku|cuhk|polyu|hkust|cityu|lingnan|hkbu|eduhk|hkmu|hksyu|thei)\b|universit(y|é) of",
                 re.I)


def scrub(path, n_slides):
    core = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
            'xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" '
            'xmlns:dcmitype="http://purl.org/dc/dcmitype/" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
            '<dc:title>Built to Last</dc:title><dc:subject>UFP Award 2026, Round Two: a financial proposal for the Wong '
            'family</dc:subject><dc:creator>Team Axis</dc:creator><cp:lastModifiedBy>Team Axis</cp:lastModifiedBy>'
            '<cp:revision>1</cp:revision><dcterms:created xsi:type="dcterms:W3CDTF">{t}</dcterms:created>'
            '<dcterms:modified xsi:type="dcterms:W3CDTF">{t}</dcterms:modified></cp:coreProperties>')
    t = datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    app = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
           '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
           'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
           '<TotalTime>0</TotalTime><Application>Microsoft Office PowerPoint</Application>'
           f'<PresentationFormat>Widescreen</PresentationFormat><Slides>{n_slides}</Slides><Notes>{n_slides}</Notes>'
           '<HiddenSlides>0</HiddenSlides><MMClips>0</MMClips><ScaleCrop>false</ScaleCrop><Company></Company>'
           '<LinksUpToDate>false</LinksUpToDate><SharedDoc>false</SharedDoc><HyperlinksChanged>false</HyperlinksChanged>'
           '<AppVersion>16.0000</AppVersion></Properties>')
    tmp = path + ".tmp"

    def target(rels_name, tgt):
        """the part a relationship points at, as a package path"""
        if tgt.startswith("/"):
            return tgt[1:]
        src_dir = os.path.dirname(os.path.dirname(rels_name))       # 'ppt/_rels/x.rels' -> 'ppt'
        return os.path.normpath(os.path.join(src_dir, tgt)).replace("\\", "/")

    with zipfile.ZipFile(path) as zin:
        names = zin.namelist()
        dropped = [n for n in names if DROP.match(n) or (("/_rels/" in n) and DROP.match(n.replace("_rels/", "")[:-5]))]
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                n = item.filename
                if n in dropped:
                    continue
                data = zin.read(n)
                if n == "docProps/core.xml":
                    data = core.replace("{t}", t).encode("utf-8")
                elif n == "docProps/app.xml":
                    data = app.encode("utf-8")
                elif n.endswith(".rels"):
                    txt = data.decode("utf-8")
                    txt = re.sub(r'<Relationship [^>]*?Target="([^"]+)"[^>]*/>',
                                 lambda m: "" if target(n, m.group(1)) in dropped else m.group(0), txt)
                    data = txt.encode("utf-8")
                elif n == "[Content_Types].xml":
                    txt = data.decode("utf-8")
                    txt = re.sub(r'<Override [^>]*?PartName="/([^"]+)"[^>]*/>',
                                 lambda m: "" if m.group(1) in dropped else m.group(0), txt)
                    data = txt.encode("utf-8")
                zout.writestr(item, data)
    os.replace(tmp, path)
    hits = []
    with zipfile.ZipFile(path) as z:
        for n in z.namelist():
            if n.endswith((".xml", ".rels", ".svg")):
                txt = z.read(n).decode("utf-8", "replace")
                txt = re.sub(r'xmlns(:\w+)?="[^"]*"', "", txt)
                hits += [(n, m.group(0)) for m in BAD.finditer(txt)]
    return dropped, hits


def check(path):
    with zipfile.ZipFile(path) as z:
        for name in ("docProps/core.xml", "docProps/app.xml"):
            if name in z.namelist():
                print(f"--- {name}\n{z.read(name).decode('utf-8')}")
        print("fonts embedded:", [n for n in z.namelist() if n.startswith("ppt/fonts/")] or "none")
        print("other metadata parts:", [n for n in z.namelist() if DROP.match(n)] or "none")


def main():
    if "--check" in sys.argv:
        check(sys.argv[sys.argv.index("--check") + 1] if len(sys.argv) > sys.argv.index("--check") + 1 else OUT)
        return
    cal = float(sys.argv[sys.argv.index("--cal") + 1]) if "--cal" in sys.argv else 0.06
    os.makedirs(WORK, exist_ok=True)
    deck = open(DECK, encoding="utf-8").read()
    if "{{" in deck or "[[" in deck:
        sys.exit("Axis.html still has unfilled placeholders: run make.py first")
    data = extract(deck)
    print("extracted:", len(data["slides"]), "slides,", len(data["assets"]), "pictures and drawings;",
          data["check"]["pseudos"], "pseudo-elements made real,", data["check"]["moved"], "elements moved by it")
    info = slide_info(deck)
    # two files: the submitted one (every slide complete) and the one shared in Zoom (one click per build step)
    only = sys.argv[sys.argv.index("--only") + 1] if "--only" in sys.argv else ""    # --only submit | present
    for present, out in ((False, OUT), (True, OUT_PRESENT)):
        tag = "present" if present else "submit"
        if only and only != tag:
            continue
        raw = os.path.join(WORK, f"Axis_raw_{tag}.pptx")
        counts = build(data, info, raw, cal, present=present)
        print(f"built ({tag}):", " ".join(f"{i}:{c}" for i, c in enumerate(counts, 1)), "shapes")
        if "--no-office" in sys.argv:
            shutil.copy(raw, out)
            print("->", out, "(no PowerPoint step: font not embedded)")
            continue
        if not present:
            for f in os.listdir(WORK):
                if f.startswith("ppt-") and f.endswith(".png"):
                    os.remove(os.path.join(WORK, f))
        tmp_out = os.path.join(WORK, f"Axis_office_{tag}.pptx")
        if os.path.exists(tmp_out):
            os.remove(tmp_out)
        print(office(raw, tmp_out, "" if present else WORK))
        dropped, hits = scrub(tmp_out, len(info))
        if hits:
            sys.exit(f"STOP: identifying text in the file: {hits[:5]}")
        try:
            shutil.copy(tmp_out, out)
        except PermissionError:
            sys.exit(f"STOP: {os.path.basename(out)} is open in PowerPoint; close it and run again "
                     f"(the new file is ready at {tmp_out})")
        print(f"-> {out} ({os.path.getsize(out) / 1e6:.1f} MB); metadata parts dropped: {dropped or 'none'}")
    if not only or only == "submit":
        check(OUT)


if __name__ == "__main__":
    main()
