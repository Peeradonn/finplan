"""
Export the assembled proposal to PowerPoint, page for page (A4 portrait slides, editable text).
Usage:  python assemble.py && python export_pptx.py      -> build/proposal.pptx

How it works: Chrome lays out proposal.html exactly as it prints; a script in the page reads back the
position and style of every text block, image, fill and rule, and this file places each one as a native
PowerPoint shape at the same position. Text stays editable; charts are the same PNGs as the PDF.
Fonts: Times New Roman, and EB Garamond (install document/fonts/*.ttf on any machine that edits the file).
"""
import os, re, sys, json, html, subprocess, urllib.parse
from pptx import Presentation
from pptx.util import Mm, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from PIL import Image, ImageEnhance

import fill

HERE = os.path.dirname(os.path.abspath(__file__))
BUILD = os.path.join(HERE, "build")
ASSETS = os.path.join(BUILD, "pptx-assets")
os.makedirs(ASSETS, exist_ok=True)
SRC = os.path.join(HERE, "proposal.html")
TMP = os.path.join(HERE, "_export_proposal.html")
OUT = os.path.join(BUILD, "proposal.pptx")
WIDTH_SLACK_MM = 0.8          # PowerPoint wraps a hair earlier than Chrome; this keeps line breaks identical

EXTRACT = r"""
<script>
(async () => {
await document.fonts.ready;
const MM = 96 / 25.4;
const rgba = c => { const m = c && c.match(/rgba?\(([^)]+)\)/); if (!m) return null;
  const p = m[1].split(',').map(parseFloat); const a = p.length > 3 ? p[3] : 1;
  return a === 0 ? null : [p[0], p[1], p[2], a]; };
const style = s => ({ font: s.fontFamily, size: parseFloat(s.fontSize) * 0.75, weight: parseInt(s.fontWeight),
  italic: s.fontStyle === 'italic', color: rgba(s.color),
  ls: s.letterSpacing === 'normal' ? 0 : parseFloat(s.letterSpacing) * 0.75 });
function textNodes(e) { const a = []; const w = document.createTreeWalker(e, NodeFilter.SHOW_TEXT); let t;
  while ((t = w.nextNode())) if (t.textContent.trim()) a.push(t); return a; }
function runs(el) {
  // inline margins (a label's gap before the title) become an en space, about 2mm at 12pt
  const after = new Set(), before = new Set();
  for (const e of el.querySelectorAll('*')) {
    const s = getComputedStyle(e); if (s.display !== 'inline') continue;
    const tn = textNodes(e); if (!tn.length) continue;
    if (parseFloat(s.marginRight) > 2) after.add(tn[tn.length - 1]);
    if (parseFloat(s.marginLeft) > 2) before.add(tn[0]);
  }
  const out = []; const tw = document.createTreeWalker(el, NodeFilter.SHOW_TEXT | NodeFilter.SHOW_ELEMENT); let n;
  while ((n = tw.nextNode())) {
    if (n.nodeType === 1) { if (n.tagName === 'BR') out.push({ br: true }); continue; }
    const s = getComputedStyle(n.parentElement);
    let t = n.textContent.replace(/[\t\n\r ]+/g, ' ');
    if (s.textTransform === 'uppercase') t = t.toUpperCase();
    if (after.has(n)) t = t.replace(/ $/, '') + ' ';
    if (before.has(n)) t = ' ' + t.replace(/^ /, '');
    if (t) out.push(Object.assign({ text: t }, style(s)));
  }
  // collapse white space the way CSS does
  let prevSpace = true;
  for (const r of out) {
    if (r.br) { prevSpace = true; continue; }
    if (prevSpace) r.text = r.text.replace(/^ /, '');
    if (r.text) prevSpace = r.text.endsWith(' ');
  }
  for (let i = out.length - 1; i >= 0; i--) { if (out[i].br) continue; out[i].text = out[i].text.replace(/ $/, ''); if (out[i].text) break; }
  for (let i = 0; i < out.length; i++) if (out[i].br && i > 0 && !out[i - 1].br) out[i - 1].text = out[i - 1].text.replace(/ $/, '');
  return out.filter(r => r.br || r.text);
}
function textItem(el, cs, pr, extra) {
  const r = el.getBoundingClientRect();
  const pl = parseFloat(cs.paddingLeft) + parseFloat(cs.borderLeftWidth), pt = parseFloat(cs.paddingTop) + parseFloat(cs.borderTopWidth);
  const prr = parseFloat(cs.paddingRight) + parseFloat(cs.borderRightWidth), pb = parseFloat(cs.paddingBottom) + parseFloat(cs.borderBottomWidth);
  const lh = cs.lineHeight === 'normal' ? parseFloat(cs.fontSize) * 1.15 : parseFloat(cs.lineHeight);
  const flex = cs.display.includes('flex') || cs.display.includes('grid');
  const align = flex && cs.justifyContent === 'center' ? 'center' : cs.textAlign;
  const middle = flex && (cs.alignItems === 'center' || cs.alignContent === 'center');
  return Object.assign({ t: 'text', x: (r.left - pr.left + pl) / MM, y: (r.top - pr.top + pt) / MM,
    w: (r.width - pl - prr) / MM, h: (r.height - pt - pb) / MM, align, middle, lh: lh * 0.75,
    vertical: cs.writingMode.startsWith('vertical'), runs: runs(el) }, extra || {});
}
function walk(el, pr, items) {
  const cs = getComputedStyle(el);
  if (cs.display === 'none' || cs.visibility === 'hidden' || ['SCRIPT', 'STYLE'].includes(el.tagName)) return;
  const r = el.getBoundingClientRect();
  const box = { x: (r.left - pr.left) / MM, y: (r.top - pr.top) / MM, w: r.width / MM, h: r.height / MM };
  const bg = rgba(cs.backgroundColor);
  if (bg && r.width > 0 && r.height > 0) items.push(Object.assign({ t: 'rect', fill: bg }, box));
  if (cs.backgroundImage && cs.backgroundImage !== 'none') {
    const m = cs.backgroundImage.match(/url\("?(.*?)"?\)/);
    if (m) items.push(Object.assign({ t: 'bgimg', src: m[1], pos: cs.backgroundPosition, filter: cs.filter }, box));
  }
  for (const side of ['Top', 'Right', 'Bottom', 'Left']) {
    const w = parseFloat(cs['border' + side + 'Width']), c = rgba(cs['border' + side + 'Color']);
    if (w > 0 && c && cs['border' + side + 'Style'] !== 'none') items.push(Object.assign({ t: 'border', side, bw: w / MM, color: c }, box));
  }
  if (el.tagName === 'IMG') {
    const nw = el.naturalWidth, nh = el.naturalHeight;
    let w = r.width, h = r.height, x = r.left, y = r.top;
    if (cs.objectFit === 'contain') { const k = Math.min(r.width / nw, r.height / nh); w = nw * k; h = nh * k; }
    items.push({ t: 'img', src: el.currentSrc || el.src, x: (x - pr.left) / MM, y: (y - pr.top) / MM, w: w / MM, h: h / MM });
    return;
  }
  // counters drawn by ::before (recommendation and pillar numbers)
  const par = el.parentElement;
  if (el.tagName === 'LI' && par && (par.classList.contains('recs') || par.classList.contains('pillars'))) {
    const b = getComputedStyle(el, '::before');
    const idx = [...par.children].indexOf(el) + 1;
    items.push({ t: 'text', x: box.x, y: box.y + (parseFloat(b.top) || 0) / MM, w: 10, h: 8, align: 'left',
      lh: parseFloat(b.fontSize) * 0.75 * 1.1, vertical: false,
      runs: [Object.assign({ text: String(idx).padStart(2, '0') }, style(b))] });
  }
  // decorative bar drawn by ::after (pull quote)
  const a = getComputedStyle(el, '::after');
  if (a.content === '""' && rgba(a.backgroundColor) && a.display === 'block') {
    const ah = parseFloat(a.height) / MM;
    items.push({ t: 'rect', fill: rgba(a.backgroundColor), x: box.x, y: box.y + box.h - ah, w: parseFloat(a.width) / MM, h: ah });
  }
  const kids = [...el.childNodes];
  const hasText = kids.some(n => n.nodeType === 3 && n.textContent.trim());
  const elKids = kids.filter(n => n.nodeType === 1 && getComputedStyle(n).display !== 'none');
  const allInline = elKids.every(k => getComputedStyle(k).display === 'inline');
  if (cs.display !== 'inline' && el.textContent.trim() && allInline) { items.push(textItem(el, cs, pr)); return; }
  if (hasText) console.warn('mixed text and blocks', el.className);
  for (const k of elKids) walk(k, pr, items);
}
const pages = [];
for (const p of document.querySelectorAll('section.page')) {
  const pr = p.getBoundingClientRect(); const items = [];
  walk(p, pr, items);
  pages.push({ w: pr.width / MM, h: pr.height / MM, items });
}
const pre = document.createElement('pre'); pre.id = '__export'; pre.style.display = 'none';
pre.textContent = JSON.stringify(pages); document.body.appendChild(pre);
})();
</script>
"""

def extract():
    with open(SRC, encoding="utf-8") as f:
        page = fill.fill_text(f.read())
    page = re.sub(r"<script>.*?</script>", "", page, flags=re.S).replace("</body>", EXTRACT + "</body>")
    with open(TMP, "w", encoding="utf-8") as f:
        f.write(page)
    chrome = next(b for b in (r"C:\Program Files\Google\Chrome\Application\chrome.exe",
                              r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe") if os.path.exists(b))
    dom = subprocess.run([chrome, "--headless=new", "--disable-gpu", "--no-first-run", "--virtual-time-budget=15000",
                          "--dump-dom", "file:///" + TMP.replace("\\", "/")],
                         check=True, capture_output=True, text=True, encoding="utf-8").stdout
    os.remove(TMP)
    m = re.search(r'<pre id="__export"[^>]*>(.*?)</pre>', dom, re.S) or sys.exit("layout export failed")
    return json.loads(html.unescape(m.group(1)))

def path_of(url):
    return urllib.parse.unquote(urllib.parse.urlparse(url).path).lstrip("/")

def rgb(c):
    return RGBColor(*(int(round(v)) for v in c[:3]))

def font_name(family, weight):
    fam = family.split(",")[0].strip().strip('"')
    if "Garamond" in fam:
        return {500: "EB Garamond Medium", 600: "EB Garamond SemiBold"}.get(weight, "EB Garamond"), False
    return fam, weight >= 600          # Times has no semibold: 600 prints as bold

def add_rect(slide, x, y, w, h, color):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Mm(x), Mm(y), Mm(max(w, 0.05)), Mm(max(h, 0.05)))
    s.fill.solid(); s.fill.fore_color.rgb = rgb(color); s.line.fill.background(); s.shadow.inherit = False
    if len(color) > 3 and color[3] < 1:
        clr =s.fill._xPr.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr")
        a = clr.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}alpha", {"val": str(int(color[3] * 100000))})
        clr.append(a)
    return s

def cover_image(item):
    src = path_of(item["src"])
    im = Image.open(src).convert("RGB")
    m = re.search(r"saturate\(([\d.]+)\)", item.get("filter") or "")
    if m:
        im = ImageEnhance.Color(im).enhance(float(m.group(1)))
    bw, bh = item["w"], item["h"]; iw, ih = im.size
    k = max(bw / iw, bh / ih)                               # background-size: cover
    cw, ch = bw / k, bh / k
    px, py = [float(v.strip("%")) / 100 if v.endswith("%") else 0.5 for v in item["pos"].split()[:2]]
    left, top = (iw - cw) * px, (ih - ch) * py
    im = im.crop((int(left), int(top), int(left + cw), int(top + ch)))
    out = os.path.join(ASSETS, "cover.jpg"); im.save(out, quality=92)
    return out

def add_text(slide, it):
    x, y, w, h = it["x"], it["y"], it["w"] + WIDTH_SLACK_MM, it["h"]
    if it["vertical"]:
        cx, cy = x + it["w"] / 2, y + h / 2
        w, h = it["h"] + WIDTH_SLACK_MM, it["w"]
        x, y = cx - w / 2, cy - h / 2
    tb = slide.shapes.add_textbox(Mm(x), Mm(y), Mm(w), Mm(max(h, 1)))
    if it["vertical"]:
        tb.rotation = 270
    tf = tb.text_frame
    tf.word_wrap = True; tf.auto_size = MSO_AUTO_SIZE.NONE; tf.vertical_anchor = MSO_ANCHOR.MIDDLE if it.get("middle") else MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = {"right": PP_ALIGN.RIGHT, "center": PP_ALIGN.CENTER, "justify": PP_ALIGN.JUSTIFY}.get(it["align"], PP_ALIGN.LEFT)
    p.line_spacing = Pt(it["lh"])
    for r in it["runs"]:
        if r.get("br"):
            p.add_line_break(); continue
        run = p.add_run(); run.text = r["text"]
        name, bold = font_name(r["font"], r["weight"])
        f = run.font; f.name = name; f.size = Pt(round(r["size"] * 2) / 2); f.bold = bold; f.italic = r["italic"]
        if r["color"]:
            f.color.rgb = rgb(r["color"])
        if r["ls"]:
            run._r.get_or_add_rPr().set("spc", str(int(round(r["ls"] * 100))))
    return tb

def build(pages):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Mm(pages[0]["w"]), Mm(pages[0]["h"])
    prs.core_properties.title = "Built to Last: A Financial Plan for the Wong Family"
    prs.core_properties.author = "Team Axis"
    prs.core_properties.last_modified_by = "Team Axis"
    blank = prs.slide_layouts[6]
    counts = {}
    for pg in pages:
        slide = prs.slides.add_slide(blank)
        for it in pg["items"]:
            counts[it["t"]] = counts.get(it["t"], 0) + 1
            if it["t"] == "rect":
                add_rect(slide, it["x"], it["y"], it["w"], it["h"], it["fill"])
            elif it["t"] == "border":
                s, bw = it["side"], it["bw"]
                x, y, w, h = it["x"], it["y"], it["w"], it["h"]
                geo = {"Top": (x, y, w, bw), "Bottom": (x, y + h - bw, w, bw),
                       "Left": (x, y, bw, h), "Right": (x + w - bw, y, bw, h)}[s]
                add_rect(slide, *geo, it["color"])
            elif it["t"] == "bgimg":
                slide.shapes.add_picture(cover_image(it), Mm(it["x"]), Mm(it["y"]), Mm(it["w"]), Mm(it["h"]))
            elif it["t"] == "img":
                slide.shapes.add_picture(path_of(it["src"]), Mm(it["x"]), Mm(it["y"]), Mm(it["w"]), Mm(it["h"]))
            elif it["t"] == "text":
                add_text(slide, it)
    prs.save(OUT)
    return counts

EMBED_PS = r"""
$pp = New-Object -ComObject PowerPoint.Application
try {
  $pres = $pp.Presentations.Open('%s', $true, $false, $false)
  $pres.RemovePersonalInformation = -1      # every later save strips author and editor names too
  $pres.SaveAs('%s', 24, -1)                # 24 = .pptx; -1 = embed TrueType fonts
  $pres.Close()
} finally { $pp.Quit() }
"""

def embed_fonts():
    """Re-save through PowerPoint (if installed) so EB Garamond travels inside the file."""
    tmp_in, tmp_out = OUT + ".plain.pptx", OUT + ".embedded.pptx"
    os.replace(OUT, tmp_in)
    ps = os.path.join(BUILD, "_embed.ps1")
    with open(ps, "w", encoding="utf-8") as f:
        f.write(EMBED_PS % (tmp_in, tmp_out))
    r = subprocess.run(["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", ps], capture_output=True, text=True)
    os.remove(ps)
    if r.returncode or not os.path.exists(tmp_out):
        os.replace(tmp_in, OUT)
        print("PowerPoint not available: fonts not embedded (install document/fonts to edit).")
        return False
    os.remove(tmp_in); os.replace(tmp_out, OUT)
    return True

def scrub_metadata():
    """Office writes the signed-in account into lastModifiedBy; the rules forbid any university name in
    file metadata, so rewrite the core properties and check every part for leftovers."""
    import zipfile, datetime
    now = datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    tmp = OUT + ".tmp"
    with zipfile.ZipFile(OUT) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "docProps/core.xml":
                x = data.decode("utf-8")
                x = re.sub(r"<cp:lastModifiedBy>.*?</cp:lastModifiedBy>", "<cp:lastModifiedBy>Team Axis</cp:lastModifiedBy>", x)
                x = re.sub(r"<dc:creator>.*?</dc:creator>", "<dc:creator>Team Axis</dc:creator>", x)
                x = re.sub(r"<dc:description>.*?</dc:description>", "<dc:description></dc:description>", x)
                x = re.sub(r'(<dcterms:created xsi:type="dcterms:W3CDTF">).*?(</dcterms:created>)', r"\g<1>" + now + r"\g<2>", x)
                data = x.encode("utf-8")
            zout.writestr(item, data)
    os.replace(tmp, OUT)
    leaks = []
    with zipfile.ZipFile(OUT) as z:
        for n in z.namelist():
            if n.endswith((".xml", ".rels")):
                t = z.read(n).decode("utf-8", "ignore").lower()
                words = ("hku", "universit", "college", "@") if n.startswith("docProps/") else ("hku",)
                leaks += [f"{n}: {w}" for w in words if w in t]
    return leaks

if __name__ == "__main__":
    pages = extract()
    counts = build(pages)
    embedded = embed_fonts()
    leaks = scrub_metadata()
    print(f"{OUT}: {len(pages)} slides, A4 portrait; fonts embedded: {embedded}; shapes by kind: {counts}")
    if leaks:
        sys.exit("Identifying text left in the file: " + "; ".join(leaks))
    print("metadata: author and last editor = Team Axis; no university or account name in any part")
